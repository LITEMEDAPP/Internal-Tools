import os
from openpyxl import Workbook, load_workbook
from config import test_steps

def save_results_to_excel(all_results, filename):

    if os.path.exists(filename):
        # --- Load workbook ---
        wb_existing = load_workbook(filename)
        ws_existing = wb_existing.active

        # --- Check if timestamps exist ---
        has_timestamps = (ws_existing.cell(row=1, column=3).value == "Start Time")

        # Add new columns to old format files
        if not has_timestamps:
            ws_existing.insert_cols(3)  # Insert Start Time column
            ws_existing.cell(row=1, column=3, value="Start Time")

            step_count = len(test_steps.keys())
            end_time_col = 3 + step_count + 1

            ws_existing.cell(row=1, column=end_time_col, value="End Time")
            ws_existing.cell(row=1, column=end_time_col + 1, value="Battery Value")
            ws_existing.cell(row=1, column=end_time_col + 2, value="Dose read response")
            ws_existing.cell(row=1, column=end_time_col + 3, value="Device UID")
            ws_existing.cell(row=1, column=end_time_col + 4, value="HW FW version")

            for row in range(2, ws_existing.max_row + 1):
                ws_existing.cell(row=row, column=3, value="N/A")
                ws_existing.cell(row=row, column=end_time_col, value="N/A")
                ws_existing.cell(row=row, column=end_time_col + 1, value="")
                ws_existing.cell(row=row, column=end_time_col + 2, value="")
                ws_existing.cell(row=row, column=end_time_col + 3, value="")
                ws_existing.cell(row=row, column=end_time_col + 4, value="")

        # --- Existing rows map ---
        existing_rows = {ws_existing.cell(row=i, column=1).value: i
                         for i in range(2, ws_existing.max_row + 1)}

        # --- Update or append rows ---
        for device_name, info in all_results.items():

            if not isinstance(info, dict):
                print(f"⚠ Skipping invalid result for {device_name}")
                continue

            row_data = [
                device_name,
                info.get("address", ""),
                info.get("start_time", "N/A"),
                *info.get("steps", []),
                info.get("end_time", "N/A"),
                info.get("get_bat_response"),
                info.get("dose_set_response", ""),
                info.get("dev_uid_response", ""),
                info.get("dev_HW_FW_version", "")
            ]

            if device_name in existing_rows:
                row = existing_rows[device_name]
                for col, value in enumerate(row_data, start=1):
                    ws_existing.cell(row=row, column=col).value = value
            else:
                ws_existing.append(row_data)

        wb_existing.save(filename)

    else:
        # --- Create new workbook ---
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"

        header = [
            "Device Name",
            "MAC Address",
            "Start Time",
            *test_steps.keys(),
            "End Time",
            "Battery (%)",
            "Dose ",
            "M UID",
            "HW FW version"
        ]
        ws.append(header)

        # --- Add data rows ---
        for device_name, info in all_results.items():

            if not isinstance(info, dict):
                print(f"⚠ Skipping invalid result for {device_name}")
                continue

            ws.append([
                device_name,
                info.get("address", ""),
                info.get("start_time", "N/A"),
                *info.get("steps", []),
                info.get("end_time", "N/A"),
                info.get("get_bat_response"),
                info.get("dose_set_response", ""),
                info.get("dev_uid_response", ""),
                info.get("dev_HW_FW_version", "")
            ])

        wb.save(filename)
